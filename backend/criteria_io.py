"""XLSX import / export for the criteria configuration.

Columns (case-insensitive headers):
    name        — short criterion name
    description — full description (rubric text)
    weight      — numeric weight (e.g. 1, 2, 0.5)
    max_score   — max score (typically 5)
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook


REQUIRED_COLUMNS = ("name", "description", "weight", "max_score")


def criteria_to_xlsx_bytes(criteria: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "criteria"
    ws.append(list(REQUIRED_COLUMNS))
    for c in criteria:
        ws.append(
            [
                str(c.get("name", "")),
                str(c.get("description", "")),
                float(c.get("weight", 1.0)),
                float(c.get("max_score", 5.0)),
            ]
        )
    # Reasonable column widths
    widths = {"A": 40, "B": 70, "C": 10, "D": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_to_criteria(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Spreadsheet is empty.")

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Expected headers: {', '.join(REQUIRED_COLUMNS)}."
        )
    idx = {col: header.index(col) for col in REQUIRED_COLUMNS}

    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not row or all(c in (None, "") for c in row):
            continue
        name = str(row[idx["name"]] or "").strip()
        if not name:
            continue
        description = str(row[idx["description"]] or "").strip()
        try:
            weight = float(row[idx["weight"]]) if row[idx["weight"]] is not None else 1.0
        except (TypeError, ValueError):
            raise ValueError(f"Row {len(out) + 2}: 'weight' must be numeric.")
        try:
            max_score = (
                float(row[idx["max_score"]])
                if row[idx["max_score"]] is not None
                else 5.0
            )
        except (TypeError, ValueError):
            raise ValueError(f"Row {len(out) + 2}: 'max_score' must be numeric.")
        out.append(
            {
                "name": name,
                "description": description,
                "weight": weight,
                "max_score": max_score,
            }
        )
    if not out:
        raise ValueError("No criteria rows found in the spreadsheet.")
    return out
